"""Reading raw Excel/CSV exports and turning chosen columns into scan jobs.

Customs and trade exports have one row per shipment, so the same importer
appears dozens of times. Collapsing those into one job per company is the
difference between 1007 lookups and 427.
"""

import csv
import io
import os
import re
from collections import Counter, OrderedDict
from dataclasses import dataclass, field

MAX_PREVIEW_ROWS = 8
EXCEL_EXTS = (".xlsx", ".xlsm")
CSV_EXTS = (".csv", ".tsv", ".txt")


@dataclass
class SheetInfo:
    name: str
    rows: int
    columns: int


@dataclass
class TablePreview:
    sheets: list[SheetInfo]
    sheet: str
    headers: list[str]
    sample: list[list[str]]
    total_rows: int


@dataclass
class ScanJob:
    company: str
    country: str | None = None
    products: list[str] = field(default_factory=list)
    occurrences: int = 1

    def as_row(self) -> tuple[str, str | None, list[str]]:
        return (self.company, self.country, self.products)


class UnsupportedFile(ValueError):
    pass


def _ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# --- Loading ---------------------------------------------------------------


def _load_csv(path: str) -> tuple[list[str], list[list[str]]]:
    with open(path, "r", encoding="utf-8-sig", errors="replace", newline="") as f:
        head = f.read(64_000)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(head, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        rows = [[_cell(c) for c in row] for row in reader]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def _load_excel(path: str, sheet: str | None) -> tuple[list[str], list[list[str]], list[SheetInfo], str]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [
            SheetInfo(name=n, rows=(wb[n].max_row or 0), columns=(wb[n].max_column or 0))
            for n in wb.sheetnames
        ]
        active = sheet if sheet in wb.sheetnames else wb.sheetnames[0]
        ws = wb[active]

        headers: list[str] = []
        body: list[list[str]] = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            cells = [_cell(c) for c in row]
            if i == 0:
                headers = cells
            elif any(cells):
                body.append(cells)
        return headers, body, sheets, active
    finally:
        wb.close()


def load_table(path: str, sheet: str | None = None):
    """Return (headers, rows, sheets, active_sheet) for a spreadsheet or CSV."""
    ext = _ext(path)
    if ext in EXCEL_EXTS:
        return _load_excel(path, sheet)
    if ext in CSV_EXTS:
        headers, rows = _load_csv(path)
        name = os.path.basename(path)
        return headers, rows, [SheetInfo(name=name, rows=len(rows) + 1, columns=len(headers))], name
    if ext == ".xls":
        raise UnsupportedFile(
            "Legacy .xls is not supported — re-save the file as .xlsx and try again."
        )
    raise UnsupportedFile(f"Unsupported file type: {ext or 'unknown'}")


def preview_table(path: str, sheet: str | None = None) -> TablePreview:
    headers, rows, sheets, active = load_table(path, sheet)
    # De-duplicate blank/repeated headers so the UI can key on them safely
    seen: Counter = Counter()
    clean_headers = []
    for i, h in enumerate(headers):
        name = h or f"Column {i + 1}"
        seen[name] += 1
        clean_headers.append(name if seen[name] == 1 else f"{name} ({seen[name]})")

    return TablePreview(
        sheets=sheets,
        sheet=active,
        headers=clean_headers,
        sample=[r[: len(clean_headers)] for r in rows[:MAX_PREVIEW_ROWS]],
        total_rows=len(rows),
    )


# --- Column guessing -------------------------------------------------------

_COMPANY_HINTS = [
    "tên cty nhập khẩu", "tên công ty nhập khẩu", "importer", "buyer",
    "consignee", "customer", "tên cty", "công ty", "company", "cty",
    "supplier", "exporter", "tên khách hàng", "name",
]
_COUNTRY_HINTS = [
    "tên nước nhập khẩu", "nước nhập khẩu", "importer country", "country",
    "quốc gia", "nước", "destination", "nation",
]
_PRODUCT_HINTS = [
    "tên hàng", "mặt hàng", "sản phẩm", "product", "goods",
    "description", "commodity", "item", "hs code",
]


def _guess(headers: list[str], hints: list[str]) -> str | None:
    lowered = [(h, h.lower().strip()) for h in headers]
    for hint in hints:  # hints are ordered most- to least-specific
        for original, low in lowered:
            if low == hint:
                return original
        for original, low in lowered:
            if hint in low:
                return original
    return None


def guess_columns(headers: list[str]) -> dict[str, str | None]:
    return {
        "company": _guess(headers, _COMPANY_HINTS),
        "country": _guess(headers, _COUNTRY_HINTS),
        "product": _guess(headers, _PRODUCT_HINTS),
    }


# --- Product terms ---------------------------------------------------------

_PRODUCT_SPLIT = re.compile(r"[,;(){}\[\]#/|]")
_PRODUCT_NOISE = re.compile(r"\b(hàng mới|100%|nsx|sp|tên kh|tên khoa học|dùng để|đóng gói)\b", re.I)


def sanitize_product(text: str, max_words: int = 6) -> str:
    """Trim a shipment description down to something worth searching for.

    Customs `Tên hàng` values are full descriptions — "Vỏ quế khô cắt đoạn
    dài <50 cm, chưa qua chế biến (SP thu mua...)". Pasting that into a search
    engine finds nothing, so keep only the leading clause.
    """
    if not text:
        return ""
    first = _PRODUCT_SPLIT.split(text)[0]
    first = _PRODUCT_NOISE.sub(" ", first)
    first = re.sub(r"[<>]+\s*\d*\s*\w*", " ", first)
    words = [w for w in first.split() if w]
    trimmed = " ".join(words[:max_words]).strip(" .-–:")
    return trimmed if len(trimmed) >= 3 else ""


# --- Job building ----------------------------------------------------------


def build_jobs(
    headers: list[str],
    rows: list[list[str]],
    company_col: str,
    country_col: str | None = None,
    product_col: str | None = None,
    dedupe: bool = True,
    limit: int | None = None,
    max_products: int = 3,
    product_override: list[str] | None = None,
) -> list[ScanJob]:
    """Collapse raw rows into one scan job per company.

    `product_override` replaces the per-row product terms entirely — useful
    when the sheet's description column is in another language and you know
    the English keywords for the whole file.
    """
    if company_col not in headers:
        raise ValueError(f"Column not found: {company_col}")

    ci = headers.index(company_col)
    coi = headers.index(country_col) if country_col in headers else None
    pi = headers.index(product_col) if product_col in headers else None

    def at(row: list[str], index: int | None) -> str:
        if index is None or index >= len(row):
            return ""
        return (row[index] or "").strip()

    override = [p.strip() for p in (product_override or []) if p and p.strip()]

    def products_for(raw: str) -> list[str]:
        if override:
            return list(override)
        term = sanitize_product(raw)
        return [term] if term else []

    if not dedupe:
        jobs = []
        for row in rows:
            company = at(row, ci)
            if not company:
                continue
            jobs.append(ScanJob(
                company=company,
                country=at(row, coi) or None,
                products=products_for(at(row, pi)),
            ))
        return jobs[:limit] if limit else jobs

    grouped: "OrderedDict[str, dict]" = OrderedDict()
    for row in rows:
        company = at(row, ci)
        if not company:
            continue
        key = " ".join(company.lower().split())
        bucket = grouped.setdefault(
            key,
            {"name": Counter(), "country": Counter(), "products": Counter(), "count": 0},
        )
        bucket["count"] += 1
        bucket["name"][company] += 1
        country = at(row, coi)
        if country:
            bucket["country"][country] += 1
        product = sanitize_product(at(row, pi))
        if product:
            bucket["products"][product] += 1

    jobs = []
    for bucket in grouped.values():
        # Products are free-text shipment descriptions; the most frequent few
        # are the useful search terms, the rest is noise.
        products = (
            list(override)
            if override
            else [p for p, _ in bucket["products"].most_common(max_products)]
        )
        jobs.append(ScanJob(
            company=bucket["name"].most_common(1)[0][0],
            country=(bucket["country"].most_common(1)[0][0] if bucket["country"] else None),
            products=products,
            occurrences=bucket["count"],
        ))

    # Frequent trading partners first — they are the leads worth having
    jobs.sort(key=lambda j: -j.occurrences)
    return jobs[:limit] if limit else jobs


def jobs_to_csv(jobs: list[ScanJob]) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["company", "country", "product", "occurrences"])
    for job in jobs:
        writer.writerow([job.company, job.country or "", ", ".join(job.products), job.occurrences])
    return buf.getvalue()
